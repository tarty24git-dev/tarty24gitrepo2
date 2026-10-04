import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pptx import Presentation
from pptx.util import Inches as PptInches, Pt as PptPt
from pptx.dml.color import RGBColor as PptRGBColor

def generate_word_doc():
    doc = docx.Document()

    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Title Block
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("FUNCTIONAL SPECIFICATION & WALKTHROUGH DOCUMENT\n")
    title_run.bold = True
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = RGBColor(15, 23, 42) # Slate Dark

    sub_title = title_p.add_run("CTH Enterprise Reporting System - Java 21 Spring Boot Web Platform")
    sub_title.font.size = Pt(14)
    sub_title.font.color.rgb = RGBColor(13, 110, 253)

    doc.add_paragraph()

    # Document Control Metadata Table
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    meta_data = [
        ("Project Name", "CTH Enterprise Reporting System"),
        ("Document ID", "CTH-FSW-2026-V1"),
        ("Author / Owner", "CTH Engineering & Architecture Team"),
        ("Technology Stack", "Java 21, Spring Boot 3.2.5, Maven, Thymeleaf, WebSocket, Swagger"),
        ("Target Databases", "Oracle, PostgreSQL, Microsoft SQL Server, H2 Embedded")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = table.rows[i]
        r0 = row.cells[0].paragraphs[0].add_run(k)
        r0.bold = True
        r0.font.color.rgb = RGBColor(15, 23, 42)
        row.cells[1].paragraphs[0].add_run(v)

    doc.add_page_break()

    # Section 1
    h1 = doc.add_heading("1. Executive Summary & System Overview", level=1)
    doc.add_paragraph(
        "The CTH Enterprise Reporting System is a state-of-the-art, high-performance web application engineered "
        "to deliver real-time operational analytics, dynamic database introspection, and customizable business reporting. "
        "Built on standard enterprise Java 21 and Spring Boot 3.2.5, the platform features cross-platform compatibility "
        "across Windows and Linux deployment environments."
    )

    # Section 2
    doc.add_heading("2. Functional & Technical Requirements", level=1)

    doc.add_heading("2.1 Dynamic Database Engine Switching", level=2)
    doc.add_paragraph(
        "The application features dynamic connection configuration supporting Oracle (ojdbc11), PostgreSQL, "
        "Microsoft SQL Server, and embedded H2 databases. Connections and parameters are dynamically saved and updated "
        "from the system settings UI and stored in application properties files."
    )

    doc.add_heading("2.2 Dynamic Query & Report Builder", level=2)
    doc.add_paragraph(
        "Users can inspect database metadata in real-time, select available tables and columns, set filter criteria (WHERE clauses), "
        "and define sorting directions (ASC/DESC). Interactive query results are generated online instantly with optional "
        "report configuration persistence."
    )

    doc.add_heading("2.3 Authentication & Security Model", level=2)
    doc.add_paragraph(
        "Supports hybrid authentication modes configurable directly from the frontend UI:\n"
        " • BCrypt Encrypted Local Database Authentication\n"
        " • Active Directory / LDAP Integration\n"
        " • Account maintenance including User Lock/Unlock, Password Reset, and Expiration\n"
        " • Configurable SMS/Email 2FA Challenge Toggle"
    )

    doc.add_heading("2.4 Executive Realtime Dashboard", level=2)
    doc.add_paragraph(
        "Integrates Spring WebSocket / STOMP protocol to stream real-time system metrics (CPU load, memory heap, "
        "active workflow tasks) every 3 seconds without page refreshes."
    )

    # Section 3
    doc.add_heading("3. System Walkthrough & Deployment Instructions", level=1)
    doc.add_heading("3.1 Windows Environment Deployment", level=2)
    doc.add_paragraph(
        "1. Build application package: Run 'build.bat'\n"
        "2. Start background service: Run 'start.bat'\n"
        "3. Stop running application: Run 'stop.bat'"
    )

    doc.add_heading("3.2 Linux / Unix Environment Deployment", level=2)
    doc.add_paragraph(
        "1. Build application package: Run './build.sh'\n"
        "2. Start background service: Run './start.sh'\n"
        "3. Stop running application: Run './stop.sh'"
    )

    doc.add_heading("3.3 Docker & Kubernetes Deployment", level=2)
    doc.add_paragraph(
        " • Docker Compose: 'docker-compose up -d'\n"
        " • Kubernetes: 'kubectl apply -f k8s/'"
    )

    doc.save("Functional_Specification_and_Walkthrough.docx")
    print("Generated Functional_Specification_and_Walkthrough.docx")


def generate_excel_matrix():
    wb = openpyxl.Workbook()

    # Sheet 1: Test Cases
    ws_tests = wb.active
    ws_tests.title = "JUnit Test Catalog"

    headers = ["Test ID", "Test Class", "Test Method", "Module", "Target Requirement", "Expected Result", "Status"]
    ws_tests.append(headers)

    test_data = [
        ["TC-001", "UserServiceTest", "testUserRegistrationAndLocking", "Security", "User creation & account lock/unlock", "User account status transitions between ACTIVE and LOCKED", "PASSED"],
        ["TC-002", "UserServiceTest", "testPasswordResetAndExpiration", "Security", "Password expiration & reset", "Expired password flag updates and resets cleanly", "PASSED"],
        ["TC-003", "ReportServiceTest", "testAvailableTablesAndCustomQueryExecution", "Reporting", "Dynamic table inspection & query execution", "Returns table list and executes SQL queries", "PASSED"],
        ["TC-004", "ReportServiceTest", "testReportConfigSaveAndRetrieval", "Reporting", "Save & view persistent report configs", "Report config persisted in DB and retrievable", "PASSED"],
        ["TC-005", "ControllerIntegrationTest", "testPublicLoginPage", "UI / Web", "Public login page accessibility", "Returns HTTP 200 and renders login view", "PASSED"],
        ["TC-006", "ControllerIntegrationTest", "testAuthenticatedDashboardPage", "UI / Web", "Dashboard authorization check", "Returns HTTP 200 for authenticated ADMIN user", "PASSED"],
        ["TC-007", "ControllerIntegrationTest", "testAuthenticatedReportsPage", "UI / Web", "Report builder authorization check", "Returns HTTP 200 for authenticated user", "PASSED"]
    ]

    for row in test_data:
        ws_tests.append(row)

    # Style header row
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for cell in ws_tests[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Sheet 2: Data Dictionary
    ws_dict = wb.create_sheet(title="Database Data Dictionary")
    dict_headers = ["Table Name", "Column Name", "Data Type", "Constraint", "Description"]
    ws_dict.append(dict_headers)

    dict_data = [
        ["cth_users", "id", "BIGINT", "PRIMARY KEY", "Unique user identifier"],
        ["cth_users", "username", "VARCHAR(100)", "UNIQUE, NOT NULL", "Login credential username"],
        ["cth_users", "password", "VARCHAR(255)", "NOT NULL", "BCrypt encrypted user password"],
        ["cth_users", "role", "VARCHAR(50)", "DEFAULT 'ROLE_USER'", "Security role authority"],
        ["cth_users", "locked", "BOOLEAN", "DEFAULT FALSE", "Account locked flag"],
        ["cth_users", "password_expiration_date", "TIMESTAMP", "NOT NULL", "Password expiration timestamp"],
        ["cth_report_configs", "id", "BIGINT", "PRIMARY KEY", "Unique report config identifier"],
        ["cth_report_configs", "report_name", "VARCHAR(150)", "UNIQUE, NOT NULL", "User-defined report title"],
        ["cth_report_configs", "table_name", "VARCHAR(100)", "NOT NULL", "Source database table"],
        ["cth_report_configs", "selected_fields", "VARCHAR(2000)", "NOT NULL", "Comma-separated column names"]
    ]

    for row in dict_data:
        ws_dict.append(row)

    for cell in ws_dict[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    wb.save("JUnit_Test_Cases_and_Matrix.xlsx")
    print("Generated JUnit_Test_Cases_and_Matrix.xlsx")


def generate_powerpoint():
    prs = Presentation()

    # Slide 1: Title Slide
    blank_layout = prs.slide_layouts[6]
    slide1 = prs.slides.add_slide(blank_layout)

    tx_box = slide1.shapes.add_textbox(PptInches(1), PptInches(2), PptInches(8), PptInches(3))
    tf = tx_box.text_frame
    p1 = tf.paragraphs[0]
    p1.text = "CTH Enterprise Reporting System"
    p1.font.bold = True
    p1.font.size = PptPt(36)
    p1.font.color.rgb = PptRGBColor(15, 23, 42)

    p2 = tf.add_paragraph()
    p2.text = "Architecture, Feature Walkthrough & System Specification"
    p2.font.size = PptPt(20)
    p2.font.color.rgb = PptRGBColor(13, 110, 253)

    # Slide 2: Key Features
    slide2 = prs.slides.add_slide(blank_layout)
    tx_box2 = slide2.shapes.add_textbox(PptInches(0.8), PptInches(0.8), PptInches(8.4), PptInches(5.5))
    tf2 = tx_box2.text_frame
    p_title2 = tf2.paragraphs[0]
    p_title2.text = "Key System Capabilities"
    p_title2.font.bold = True
    p_title2.font.size = PptPt(28)
    p_title2.font.color.rgb = PptRGBColor(15, 23, 42)

    bullets = [
        "Multi-Database Support: Oracle, PostgreSQL, MS SQL Server, and H2",
        "Dynamic UI Builder: Select available tables & fields for online reports",
        "Hybrid Security: BCrypt local DB or Active Directory LDAP authentication",
        "User Maintenance: Account Lock/Unlock, Password Reset, Expiration & 2FA Switch",
        "Realtime Dashboard: Live STOMP WebSockets for system CPU, Memory & Workflow metrics",
        "Cross-Platform Deployment: Standardized start/stop scripts for Windows & Linux"
    ]
    for b in bullets:
        bp = tf2.add_paragraph()
        bp.text = "• " + b
        bp.font.size = PptPt(16)
        bp.font.color.rgb = PptRGBColor(51, 65, 85)

    # Slide 3: Technology Architecture
    slide3 = prs.slides.add_slide(blank_layout)
    tx_box3 = slide3.shapes.add_textbox(PptInches(0.8), PptInches(0.8), PptInches(8.4), PptInches(5.5))
    tf3 = tx_box3.text_frame
    p_title3 = tf3.paragraphs[0]
    p_title3.text = "Technology Stack & Architecture"
    p_title3.font.bold = True
    p_title3.font.size = PptPt(28)
    p_title3.font.color.rgb = PptRGBColor(15, 23, 42)

    arch_bullets = [
        "Backend Core: Java 21, Spring Boot 3.2.5, Spring Security, Spring Data JPA",
        "API Integration: Springdoc OpenAPI 3 / Swagger UI (/swagger-ui.html)",
        "Frontend Interface: Thymeleaf, HTML5, Bootstrap 5, SockJS & Stomp.js",
        "Build System: Apache Maven (packaged executable JAR)",
        "Containerization: Dockerfile, Docker Compose, Kubernetes Deployment Manifests"
    ]
    for ab in arch_bullets:
        bp = tf3.add_paragraph()
        bp.text = "• " + ab
        bp.font.size = PptPt(16)
        bp.font.color.rgb = PptRGBColor(51, 65, 85)

    prs.save("System_Presentation_Walkthrough.pptx")
    print("Generated System_Presentation_Walkthrough.pptx")

if __name__ == "__main__":
    generate_word_doc()
    generate_excel_matrix()
    generate_powerpoint()
